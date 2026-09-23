# -*- coding: utf-8 -*-
"""Congela e compara as saidas deterministicas do app.

Uso:
    python tools/golden_master.py congelar   # grava as referencias
    python tools/golden_master.py comparar   # exit 0 se identico, 1 se divergiu

As referencias sao a prova MECANICA de nao-regressao: comparar compara SAIDA,
nunca codigo. Ler o diff e concluir "nao mudou nada" e exatamente o que a regra
anti-regressao proibe, porque o leitor ve o que espera ver.

Duas saidas sao congeladas, escolhidas por serem deterministicas de verdade:

1. **dedup** — `deduplicate()` sobre um corpus fixo de 14 itens. O corpus nao e
   decorativo: ele existe para acionar as TRES estrategias do deduplicador
   (id exato, tipo+numero, ementa fuzzy >= 0.85). Congelar um corpus que so
   aciona a primeira congelaria 1/3 do modulo e daria verde sobre o resto.
   Verificado em 2026-09-22: as 3 disparam, uma cada, 14 -> 11 unicos.

2. **planilha** — hash dos VALORES DAS CELULAS, nao dos bytes do arquivo, de
   TODAS as abas ('Normativos' e 'Diagnostico da busca', esta alimentada por
   `diagnostico_fixo.json`). Ver `_sha_planilha`, que explica por que os bytes
   nao servem e o que obriga a recongelar.
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
APP = RAIZ / "levantamento-normativos"
GOLDEN = RAIZ / "tests" / "golden"

sys.path.insert(0, str(APP))


def _carregar_entrada() -> list:
    """Le o corpus fixo e devolve objetos NOVOS a cada chamada.

    Devolver objetos novos e load-bearing, nao higiene: `deduplicator._merge()`
    altera o registro sobrevivente IN-PLACE (a propria docstring dele diz
    "Modified in-place"), mexendo em ementa, nome, source, found_by, relevancia
    e link — todas colunas da planilha. Reaproveitar a lista entre o dedup e a
    planilha congelaria a planilha sobre itens ja mutados.
    """
    from models import NormativoResult

    dados = json.loads((GOLDEN / "entrada_fixa.json").read_text(encoding="utf-8"))
    return [NormativoResult(**d) for d in dados]


def _saida_dedup(itens: list) -> list[dict]:
    from deduplicator import deduplicate

    unicos = deduplicate(itens)
    # ordem estavel: o dedup preserva a ordem de entrada, mas nao a PROMETE no
    # contrato. Ordenar por id torna a referencia imune a uma mudanca de ordem
    # que nao seja regressao.
    return sorted(
        [{"id": r.id, "nome": r.nome, "link": r.link} for r in unicos],
        key=lambda d: d["id"],
    )


def _carregar_diagnostico() -> list:
    from models import KeywordStatus
    dados = json.loads((GOLDEN / "diagnostico_fixo.json").read_text(encoding="utf-8"))
    return [KeywordStatus(**d) for d in dados]


def _sha_planilha(itens: list) -> str:
    """Hash dos VALORES das celulas, nao dos bytes do arquivo — de TODAS as abas.

    .xlsx e um ZIP: o date_time de cada membro e o docProps/core.xml carregam o
    relogio da geracao, entao o sha dos bytes crus muda a CADA execucao, com
    entrada identica (medido por tres revisores em 2026-09-16). Congelar bytes
    faria o comparador imprimir DIVERGIU sem nada ter mudado — e o risco pior
    nao e o falso vermelho, e o executor apagar a prova para destravar a task.

    load_workbook e a mesma tecnica que test_phase4.py ja usa.

    Hasheia TODAS as abas (M7 da rodada de 22/09): a aba 'Diagnostico da busca'
    e o registro de que a fonte nao respondeu; sem ela no hash, uma regressao
    ali passaria com 'golden-master OK'. O diagnostico fixo cobre: error com
    motivo e detalhe; ok parcial; empty; nao_consultada; error retentado sem
    detalhe (so error_message). `quando` e fixo para o hash ser estavel.

    ⚠ O hash da aba de diagnostico depende de models.redigir e rotulo_status e
    de excel_export.ORIGEM_LABEL, VAZIO e do titulo com `quando` fixo (rodada 3):
    mudanca INTENCIONAL em qualquer um deles = recongelar com justificativa,
    nao regressao do dedup/export.

    A funcao publica e generate_excel(results, topic, diagnostico=None, quando=None)
    (excel_export.py) — e o que app.py e test_phase4.py usam.
    """
    from excel_export import generate_excel
    from openpyxl import load_workbook

    wb = load_workbook(generate_excel(itens, topic="golden-master",
                                      diagnostico=_carregar_diagnostico(), quando="22/09/2026 00:00"))
    linhas = []
    for ws in wb.worksheets:
        linhas.append(("__aba__", ws.title))
        linhas.extend(tuple(c.value for c in linha) for linha in ws.iter_rows())
    return hashlib.sha256(repr(linhas).encode("utf-8")).hexdigest()


def _ambiente() -> str:
    """Versoes que o hash de celulas depende.

    O hash agora depende do que o LEITOR do openpyxl devolve, e o leitor coage
    tipo. Um upgrade da biblioteca pode mudar o hash sem nada ter regredido.
    Por isso a versao fica gravada junto: divergencia logo apos upgrade de
    openpyxl e RECONGELAR, nao investigar regressao.
    """
    import openpyxl

    return f"python={sys.version.split()[0]}\nopenpyxl={openpyxl.__version__}\n"


def congelar() -> None:
    GOLDEN.mkdir(parents=True, exist_ok=True)
    (GOLDEN / "dedup_esperado.json").write_text(
        json.dumps(_saida_dedup(_carregar_entrada()), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    # entrada limpa: o dedup muta os itens in-place (deduplicator._merge)
    (GOLDEN / "planilha_sha256.txt").write_text(
        _sha_planilha(_carregar_entrada()), encoding="utf-8"
    )
    (GOLDEN / "ambiente.txt").write_text(_ambiente(), encoding="utf-8")
    print("congelado em", GOLDEN)


def comparar() -> list[str]:
    divergencias: list[str] = []

    esperado = json.loads((GOLDEN / "dedup_esperado.json").read_text(encoding="utf-8"))
    obtido = _saida_dedup(_carregar_entrada())
    if obtido != esperado:
        # contagem sozinha nao basta: duas listas do MESMO tamanho com conteudo
        # diferente produziriam "esperado 11, obtido 11", que nao diz nada.
        detalhe = ""
        for i, (e, o) in enumerate(zip(esperado, obtido)):
            if e != o:
                detalhe = f"; 1o item divergente (indice {i}): esperado {e} != obtido {o}"
                break
        divergencias.append(
            f"dedup divergiu: esperado {len(esperado)} unicos, obtido {len(obtido)}{detalhe}"
        )

    # entrada limpa: o dedup muta os itens in-place (deduplicator._merge)
    sha_esperado = (GOLDEN / "planilha_sha256.txt").read_text(encoding="utf-8").strip()
    sha_obtido = _sha_planilha(_carregar_entrada())
    if sha_obtido != sha_esperado:
        divergencias.append(f"planilha divergiu: {sha_esperado[:12]} -> {sha_obtido[:12]}")

    return divergencias


if __name__ == "__main__":
    acao = sys.argv[1] if len(sys.argv) > 1 else "comparar"
    if acao == "congelar":
        congelar()
        sys.exit(0)
    problemas = comparar()
    for p in problemas:
        print("DIVERGIU:", p)
    if problemas:
        congelado = (GOLDEN / "ambiente.txt")
        if congelado.exists():
            print("\nAmbiente do congelamento:")
            print(congelado.read_text(encoding="utf-8").strip())
            print("Agora:", _ambiente().strip().replace("\n", " | "))
            print("Se a unica diferenca for a versao de openpyxl, e RECONGELAR, nao regressao.")
    print("golden-master OK" if not problemas else f"{len(problemas)} divergencia(s)")
    sys.exit(1 if problemas else 0)
